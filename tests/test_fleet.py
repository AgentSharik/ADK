"""Парк ПК и служебные режимы: «Внимание», WoL/MAC, массовый пинг, ПО и входы, сравнение ПК, шаблоны,
уведомления, PostgreSQL-адаптер, portable, CLI и серверный режим, роли, опись ПК, безопасность фильтров/журнала.

Как читать: каждый тест — «дано → действие → проверка». Сеть и Windows не нужны — ответы подменяются готовыми данными.
"""
import json
import os
import pytest
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

from adk import attention, cli, config, db, logons, nettools, netutils, notify, pgadapter, software, templates, workers
from adk.fleet import compare_lists, compare_specs, flatten_specs
from adk.extras import password_card_text


class _Attr:
    def __init__(self, v):
        self.values = v if isinstance(v, list) else [v]
        self.value = self.values[0] if self.values else None


class _Entry:
    """Поддельная запись ldap3: достаточно, чтобы работали get_ad_value / get_ad_datetime."""

    def __init__(self, dn="CN=x,OU=Users,DC=example,DC=local", **attrs):
        self.entry_dn = dn
        self._a = {k: _Attr(v) for k, v in attrs.items()}

    def __getitem__(self, k):
        return self._a[k]


NOW = datetime(2026, 9, 5, 12, 0, tzinfo=timezone.utc)


def _user(login, **kw):
    base = dict(sAMAccountName=login, sn=login.title(), givenName="Тест", userAccountControl=512)
    base.update(kw)
    return _Entry(f"CN={login},OU=Users,DC=example,DC=local", **base)


def _spin(app, ms=300):
    import time
    t0 = time.time()
    while (time.time() - t0) * 1000 < ms:
        app.processEvents()
        time.sleep(0.01)


class _NoConn:
    entries = []
    result = {"controls": {}}

    def search(self, *a, **k):
        return True

    def unbind(self):
        return True


# ------------------------------------------------------------------ 1. «Внимание»
def test_attention_from_entries_classifies(monkeypatch):
    # сроки паролей в сводку НЕ попадают (смарт-карты/сертификаты) — даже истёкший пароль не даёт пункта
    monkeypatch.setattr(config.settings, "max_password_age_days", 90)
    entries = [
        _user("expired", pwdLastSet=NOW - timedelta(days=100), lastLogonTimestamp=NOW),   # пароль истёк — не пункт
        _user("expiring_acct", pwdLastSet=NOW, lastLogonTimestamp=NOW, accountExpires=NOW + timedelta(days=3)),  # УЗ истекает
        _user("locked", pwdLastSet=NOW, lockoutTime=NOW - timedelta(hours=1)),  # заблокирован
        _user("sleepy", pwdLastSet=NOW, lastLogonTimestamp=NOW - timedelta(days=120)),  # нет входа 120 дн.
        _user("fine", pwdLastSet=NOW, lastLogonTimestamp=NOW),                  # всё хорошо — не попадает
        _user("off", pwdLastSet=NOW - timedelta(days=100), userAccountControl=514),     # отключённая — пропускаем
    ]
    items = attention.from_entries(entries, now=NOW, acct_days=7, no_logon_days=90)
    by = {(i["kind"], i["subject"]): i for i in items}
    assert not any(k.startswith("pwd") for k, _ in by) and "expired" not in {s for _, s in by}
    assert by[("acct_expiring", "expiring_acct")]["severity"] == "medium"
    assert ("locked", "locked") in by
    assert by[("no_logon", "sleepy")]["severity"] == "low" and "120 дн." in by[("no_logon", "sleepy")]["text"]
    assert not any(i["subject"] in ("fine", "off") for i in items)
    # ключ стабилен — по нему работает «отложить»
    assert by[("locked", "locked")]["key"] == "locked:locked"


def test_attention_summary_and_snooze(monkeypatch):
    assert attention.summary_line([]) == "✅ Всё спокойно"
    items = [attention._item("locked", "a", "t", "high"), attention._item("no_logon", "b", "t", "low")]
    line = attention.summary_line(items)
    assert "2" in line and "1" in line  # всего и «срочных»
    # отложили «a» на 7 дней → collect его не покажет (проверяем через snoozed_keys)
    db.snooze("locked:a", 7, "admin")
    assert "locked:a" in db.snoozed_keys()
    assert "no_logon:b" not in db.snoozed_keys()


def test_attention_from_inventory_stale():
    db.batch_update_inventory([{"Hostname": "PC-OLD", "User": "u", "ActualIp": "10.0.0.5", "Status": "OFFLINE", "LastLogon": "—"}],
                              "2026-01-01 10:00:00")
    # last_seen_online пуст → «никогда» → попадает в список
    items = attention.from_inventory(stale_days=30)
    assert items and items[0]["subject"] == "PC-OLD" and items[0]["kind"] == "pc_stale"


# ------------------------------------------------------------------ 2. WoL, ARP, массовый пинг
def test_parse_mac_and_magic_packet():
    assert nettools.parse_mac("00-1a-2b-3c-4d-5e") == "00:1A:2B:3C:4D:5E"
    assert nettools.parse_mac("001A.2B3C.4D5E") == "00:1A:2B:3C:4D:5E"
    assert nettools.parse_mac("не мак") == "" and nettools.parse_mac("00:1A:2B") == ""
    pkt = nettools.magic_packet("00:1A:2B:3C:4D:5E")
    # magic-пакет = 6 байт 0xFF + MAC × 16
    assert len(pkt) == 102 and pkt[:6] == b"\xff" * 6 and pkt[6:12] == bytes.fromhex("001A2B3C4D5E") and pkt[-6:] == pkt[6:12]
    with pytest.raises(ValueError):
        nettools.magic_packet("мусор")


def test_parse_arp_windows_and_linux():
    win = ("Интерфейс: 10.0.0.10 --- 0xb\n"
           "  адрес в Интернете      Физический адрес      Тип\n"
           "  10.0.0.5              00-1a-2b-3c-4d-5e     динамический\n"
           "  10.0.0.55             aa-bb-cc-dd-ee-ff     динамический\n")
    assert nettools.parse_arp(win, "10.0.0.5") == "00:1A:2B:3C:4D:5E"      # не путает 10.0.0.5 и 10.0.0.55
    lin = "pc-001 (10.0.0.5) at 00:1a:2b:3c:4d:5e [ether] on eth0\n"
    assert nettools.parse_arp(lin, "10.0.0.5") == "00:1A:2B:3C:4D:5E"
    assert nettools.parse_arp(win, "10.0.0.9") == ""


def test_mac_store_roundtrip():
    db.save_mac("PC-001.example.local", "00-1a-2b-3c-4d-5e")
    assert db.get_mac("pc-001") == "00:1A:2B:3C:4D:5E"
    assert db.get_mac("PC-NONE") == ""


def test_mass_ping_uses_network_info(monkeypatch):
    calls = []

    def fake(h, use_cache=True):
        calls.append(h)
        return ("10.0.0.1", True) if h == "PC-A" else ("Не найден", False)
    monkeypatch.setattr(nettools, "get_computer_network_info", fake)
    res = dict((c, (ip, on)) for c, ip, on in nettools.mass_ping(["PC-A", "PC-B", "PC-A", ""]))
    assert res == {"PC-A": ("10.0.0.1", True), "PC-B": ("Не найден", False)}
    assert sorted(calls) == ["PC-A", "PC-B"]  # дубли и пустые не опрашиваются


def test_msg_command_shape():
    argv = nettools.msg_command("PC-001", "Перезагрузка через 5 минут", 30)
    assert argv[:2] == ["msg", "*"] and "/server:PC-001" in argv and "/time:30" in argv and argv[-1].startswith("Перезагрузка")
    with pytest.raises(ValueError):
        nettools.msg_command("PC-001", "   ")


# ------------------------------------------------------------------ 3. ПО и входы
def test_parse_software_json_dedup_and_dates():
    raw = json.dumps({"software": [
        {"name": "7-Zip", "version": "23.01", "publisher": "Igor Pavlov", "installed": "20240102"},
        {"name": "7-zip", "version": "23.01", "publisher": "Igor Pavlov", "installed": "20240102"},   # дубль без учёта регистра
        {"name": "", "version": "1"},                                                                  # без имени — мусор
        {"name": "Adobe Reader", "version": None, "publisher": None, "installed": None},
    ], "hotfixes": [{"id": "KB500", "desc": "Update", "installed": "2026-01-01"}, {"id": "KB600", "desc": "Security", "installed": "2026-03-01"}, {"desc": "no id"}]})
    d = software.parse_software_json(raw)
    assert [s["name"] for s in d["software"]] == ["7-Zip", "Adobe Reader"]
    assert d["software"][0]["installed"] == "2024-01-02"
    assert [h["id"] for h in d["hotfixes"]] == ["KB600", "KB500"]  # новые сверху


def test_software_cache_and_fleet_search():
    software.cache_software("PC-001", [{"name": "1С:Предприятие 8", "version": "8.3", "publisher": "1С", "installed": ""},
                                       {"name": "Google Chrome", "version": "128", "publisher": "Google", "installed": ""}])
    software.cache_software("PC-002", [{"name": "Google Chrome", "version": "127", "publisher": "Google", "installed": ""}])
    items, ts = software.cached_software("pc-001")
    assert len(items) == 2 and ts
    found = software.find_software("chrome")
    assert sorted(r["comp"] for r in found) == ["PC-001", "PC-002"]
    assert software.find_software("1")  == []        # короче 2 символов — не ищем
    assert software.find_software("100%_x") == []    # спецсимволы LIKE экранируются, а не ломают запрос
    assert software.software_summary()[0] == ("Google Chrome", 2)
    # повторный опрос заменяет кэш ПК целиком
    software.cache_software("PC-002", [])
    assert software.cached_software("PC-002") == ([], "")


def test_parse_events_json_filters_noise():
    raw = json.dumps([
        {"id": 4624, "ts": "2026-09-05 09:00:00", "user": "ivanov", "domain": "CORP", "type": "2", "ip": "-"},
        {"id": 4624, "ts": "2026-09-05 09:05:00", "user": "ivanov", "domain": "CORP", "type": "7", "ip": "-"},
        {"id": 4624, "ts": "2026-09-05 09:06:00", "user": "SYSTEM", "domain": "NT AUTHORITY", "type": "5"},   # служебное
        {"id": 4624, "ts": "2026-09-05 09:07:00", "user": "PC-001$", "domain": "CORP", "type": "3"},          # компьютерная учётка
        {"id": 4624, "ts": "2026-09-05 09:08:00", "user": "svc", "domain": "CORP", "type": "4"},              # batch — не показываем
        {"id": 4625, "ts": "2026-09-05 10:00:00", "user": "petrov", "domain": "CORP", "type": "10", "ip": "10.0.0.7", "status": "0xC000006A"},
    ])
    d = logons.parse_events_json(raw)
    assert d["ok"] == 2 and d["fails"] == 1
    assert d["events"][0]["kind"] == "fail" and d["events"][0]["reason"] == "неверный пароль" and d["events"][0]["type"] == "RDP"
    assert d["by_user"] == [("ivanov", 2)]
    assert d["events"][-1]["ip"] == ""  # «-» превращается в пусто
    assert logons.parse_events_json("") == {"events": [], "by_user": [], "ok": 0, "fails": 0}
    assert "успешных входов: 2" in logons.format_summary(d)


def test_get_logons_and_software_non_windows(monkeypatch):
    monkeypatch.setattr(os, "name", "posix")
    assert "error" in logons.get_logons("PC-001")
    assert "error" in software.get_software("PC-001")
    assert "Недопустимое" in logons.get_logons("PC;rm -rf")["error"]


# ------------------------------------------------------------------ 4. сравнение ПК, шаблоны, уведомления
def test_compare_specs_and_lists():
    a = {"os": {"Название": "Windows 10"}, "rams": {"1": {"Размер": "8192"}}, "cpu": {"Модель": "i5"}}
    b = {"os": {"Название": "Windows 11"}, "rams": {"1": {"Размер": "8192"}}, "cpu": {}}
    flat = flatten_specs(a)
    assert flat["ОС / Название"] == "Windows 10" and flat["ОЗУ #1 / Размер"] == "8192"
    rows = compare_specs(a, b, only_diff=True)
    assert [(r[0], r[3]) for r in rows] == [("CPU / Модель", True), ("ОС / Название", True)]
    assert compare_specs(a, b)[1] == ("ОЗУ #1 / Размер", "8192", "8192", False)
    assert flatten_specs({"error": "нет CSV"}) == {}
    lst = compare_lists(["Chrome", "7-Zip"], ["chrome", "Office"])
    assert lst == [("7-Zip", "✔", "—", True), ("Chrome", "✔", "✔", False), ("Office", "—", "✔", True)]


def test_templates_normalize_load_save(tmp_path):
    path = str(tmp_path / "templates.json")
    t = templates.normalize({"attrs": {"department": " Бухгалтерия ", "title": "", "evil": "x"},
                             "groups": ["CN=Buh,OU=g", "CN=Buh,OU=g", " "], "ou": "OU=Buh,DC=example,DC=local", "extra": 1})
    assert t == {"attrs": {"department": "Бухгалтерия"}, "groups": ["CN=Buh,OU=g"], "ou": "OU=Buh,DC=example,DC=local", "note": ""}
    templates.save({"Бухгалтер": t, "": t}, path)
    loaded = templates.load(path)
    assert list(loaded) == ["Бухгалтер"] and loaded["Бухгалтер"]["groups"] == ["CN=Buh,OU=g"]
    assert templates.load(str(tmp_path / "нет.json")) == {}
    (tmp_path / "bad.json").write_text("{не json", encoding="utf-8")
    assert templates.load(str(tmp_path / "bad.json")) == {}


def test_templates_from_entry_and_apply_groups():
    from adk import ad
    e = _user("ref", department="ИТ", memberOf=["CN=Domain Users,CN=Users,DC=x", "CN=IT,OU=g", "CN=VPN,OU=g"])
    t = templates.from_entry(e, ad.get_ad_value, ad.get_ad_list_value)
    assert t["attrs"] == {"department": "ИТ"} and t["groups"] == ["CN=IT,OU=g", "CN=VPN,OU=g"]
    assert t["ou"] == "OU=Users,DC=example,DC=local"

    class Conn:
        result = {"description": "success"}
        modified = []

        def modify(self, dn, changes):
            self.modified.append(dn)
            if dn == "CN=VPN,OU=g":
                raise RuntimeError("нет прав")
            return True
    ok, bad = templates.apply_groups(Conn(), "CN=new,OU=x", t["groups"], ad.MODIFY_ADD)
    assert ok == ["CN=IT,OU=g"] and bad == ["CN=VPN,OU=g"]


def test_notify_should_and_format(monkeypatch):
    cfg = {"smtp_host": "mail.example.local", "smtp_to": "it@example.local", "actions": ("reset_password",), "watch_logins": ("director",)}
    assert notify.should_notify("reset_password", "ivanov", cfg)
    assert not notify.should_notify("ping", "ivanov", cfg)
    assert notify.should_notify("ping", "Director", cfg)          # наблюдаемая учётка — всегда
    assert not notify.should_notify("reset_password", "ivanov", {**cfg, "smtp_host": ""})   # почта не настроена
    assert notify.should_notify("anything", "x", {**cfg, "actions": ("*",)})
    text = notify.format_event("admin", "reset_password", "ivanov", "смена при входе", db.ACTION_LABELS)
    assert "Сброс пароля" in text and "ivanov" in text and "admin" in text


def test_notify_hook_called_from_log_action(monkeypatch):
    sent = []
    monkeypatch.setattr(notify, "send_all", lambda text, cfg, subject="": sent.append(text) or [("SMTP", True)])
    monkeypatch.setattr(config.settings, "notify", {"smtp_host": "h", "smtp_to": "t", "actions": ("disable_user",), "watch_logins": ()})
    hook = lambda a, ac, t, d: notify.notify_event(a, ac, t, d, labels=db.ACTION_LABELS, background=False)  # noqa: E731
    db.AUDIT_HOOKS.append(hook)
    try:
        db.log_action("admin", "disable_user", "petrov", "")
        db.log_action("admin", "ping", "petrov", "")
    finally:
        db.AUDIT_HOOKS.remove(hook)
    assert len(sent) == 1 and "petrov" in sent[0]


def test_password_card_text():
    card = password_card_text("ivanov", "Qw3rty!_", True, "CORP")
    assert card.startswith("Логин: CORP\\ivanov") and "Qw3rty!_" in card and "сменить" in card
    assert "сменить" not in password_card_text("ivanov", "Qw3rty!_", False)


# ------------------------------------------------------------------ 5. PostgreSQL, portable
def test_pg_translate_dialect():
    t = pgadapter.translate
    assert t("SELECT * FROM t WHERE a = ? AND b = ?") == "SELECT * FROM t WHERE a = %s AND b = %s"
    assert "SERIAL PRIMARY KEY" in t("CREATE TABLE IF NOT EXISTS n (id INTEGER PRIMARY KEY AUTOINCREMENT, x TEXT)")
    up = t("INSERT OR REPLACE INTO pc_mac (computer_name, mac, ts) VALUES (?, ?, datetime('now','localtime'))")
    assert up.endswith("ON CONFLICT (computer_name) DO UPDATE SET mac = EXCLUDED.mac, ts = EXCLUDED.ts")
    assert "to_char(now(), 'YYYY-MM-DD HH24:MI:SS')" in up and "?" not in up
    assert t("INSERT OR IGNORE INTO scanned VALUES (?)").endswith("ON CONFLICT DO NOTHING")
    assert "(%s)::interval" in t("SELECT 1 WHERE x < datetime('now','localtime', ?)")
    assert "interval '-30 days'" in t("SELECT 1 WHERE x < datetime('now','localtime', '-30 days')")
    assert t("CREATE TEMP TABLE IF NOT EXISTS s (name TEXT PRIMARY KEY)").startswith("CREATE TEMPORARY TABLE")


def test_pg_backend_requires_dsn(monkeypatch):
    # backend=postgres без DSN → остаёмся на SQLite (ничего не ломается)
    monkeypatch.setattr(config.settings, "db_backend", "postgres")
    monkeypatch.setattr(config.settings, "db_dsn", "")
    assert db.db_execute_with_retry("SELECT 1", fetch="one") == (1,)


def test_portable_resolution(monkeypatch, tmp_path):
    monkeypatch.delenv("ADK_HOME", raising=False)
    monkeypatch.delenv("ADK_PORTABLE", raising=False)
    monkeypatch.setattr(config, "_app_dir", lambda: str(tmp_path))
    assert config._resolve_docs_dir().endswith(os.path.join("Documents", config.APP_NAME))
    (tmp_path / "portable").write_text("")
    assert config._resolve_docs_dir() == os.path.join(str(tmp_path), "data")
    monkeypatch.setenv("ADK_HOME", str(tmp_path / "home"))
    assert config._resolve_docs_dir() == str(tmp_path / "home")


# ------------------------------------------------------------------ 6. CLI и серверный режим
def test_cli_version_and_help(capsys):
    assert cli.main(["--version"]) == 0
    assert capsys.readouterr().out.strip().startswith("ADK ")
    assert cli.main([]) == 2  # без аргументов — справка


def test_cli_ping_and_export_inventory(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(nettools, "get_computer_network_info", lambda h, use_cache=True: ("10.0.0.1", h == "PC-A"))
    assert cli.main(["--ping", "PC-A", "PC-B", "--json"]) == 0
    out = json.loads(capsys.readouterr().out)
    assert {o["comp"]: o["online"] for o in out} == {"PC-A": True, "PC-B": False}
    db.batch_update_inventory([{"Hostname": "PC-A", "User": "u", "ActualIp": "10.0.0.1", "Status": "ACTIVE", "LastLogon": "—"}], "2026-09-05 10:00:00")
    dest = str(tmp_path / "inv.csv")
    assert cli.main(["--export-inventory", dest]) == 0
    assert "PC-A" in open(dest, encoding="utf-8-sig").read()


def test_cli_find_uses_search_worker(monkeypatch, capsys):
    from adk import ad
    import test_gui as tg
    monkeypatch.setattr(ad, "make_connection", lambda *a, **k: tg.FakeConn(tg.ENTRIES))
    monkeypatch.setattr(cli, "load_credentials", lambda: ("u", "p"), raising=False)
    assert cli.main(["--find", "иванов", "--json"]) == 0
    rows = json.loads(capsys.readouterr().out)
    assert any(r["login"] == "ivanov" for r in rows)


def test_cli_wol_unknown_mac(capsys):
    assert cli.main(["--wol", "PC-UNKNOWN"]) == 1
    assert "MAC неизвестен" in capsys.readouterr().out


def test_serve_one_iteration(monkeypatch):
    from adk.workers import PCScannerWorker
    monkeypatch.setattr(PCScannerWorker, "scan_once", classmethod(lambda cls, f, p=None: 3))
    monkeypatch.setattr(attention, "collect_from_settings", lambda f, cfg: [])
    slept = []
    assert cli.serve(1, iterations=1, sleep=slept.append) == 0
    assert slept == []  # после последней итерации не спим
    assert any(r[2] == "scan" for r in db.audit_entries())


def test_main_dispatches_to_cli(monkeypatch):
    import sys
    from adk import __main__ as m
    monkeypatch.setattr(sys, "argv", ["adk", "--version"])
    assert m.main() == 0


# ------------------------------------------------------------------ 7. диалоги и сканер
def test_v31_dialogs_smoke(qapp, monkeypatch):
    from adk.attention_ui import AttentionDialog
    from adk.extras import NotifySettingsDialog
    from adk.fleet import ComparePCDialog, LogonsDialog, MassPingDialog, SoftwareDialog
    app = SimpleNamespace(get_conn=lambda: None, admin_name="admin", searched=[])
    app.search_text = app.searched.append
    monkeypatch.setattr(nettools, "get_computer_network_info", lambda h, use_cache=True: ("10.0.0.1", False))

    items = [attention._item("locked", "ivanov", "Иванов: заблокирована", "high")]
    d = AttentionDialog(app, items=items)
    assert d.table.rowCount() == 1
    d.table.selectRow(0)
    d.snooze(7)
    assert d.table.rowCount() == 0 and "locked:ivanov" in db.snoozed_keys()
    d.close()

    d = MassPingDialog(["PC-A", "PC-B"], app)
    _spin(qapp, 800)
    assert "не в сети: 2" in d.status.text()
    d.close()

    software.cache_software("PC-A", [{"name": "Chrome", "version": "1", "publisher": "", "installed": ""}])
    d = SoftwareDialog("PC-A", app)
    assert d.table.rowCount() == 1
    d.q.setText("chr")
    d.search_fleet()
    assert d.fleet.rowCount() == 1
    d.close()

    d = LogonsDialog("PC-A", app)
    _spin(qapp)
    assert "⚠️" in d.summary.text()  # не Windows — понятная ошибка, а не падение
    d.close()

    d = ComparePCDialog("PC-A", "PC-B", app)
    _spin(qapp)
    assert d.t_soft.rowCount() == 1  # Chrome есть только на PC-A
    assert "недоступен" in d.lbl_soft.text()      # не Windows → живой опрос не прошёл → честно сказано, что данные из сохранённого
    d.close()

    d = NotifySettingsDialog(app)
    d.smtp_host.setText("mail.example.local")
    assert d.values()["smtp_host"] == "mail.example.local" and "smtp_tls" in d.values()
    d.close()


def test_register_dialog_templates(qapp, monkeypatch, tmp_path):
    from adk.dialogs import RegisterUserDialog
    path = str(tmp_path / "t.json")
    templates.save({"Бухгалтер": {"attrs": {"department": "Бухгалтерия", "title": "Бухгалтер"}, "groups": ["CN=Buh,OU=g"], "ou": "OU=Buh,DC=x"}}, path)
    monkeypatch.setattr(config.settings, "templates_file", path)
    app = SimpleNamespace(get_conn=lambda: _NoConn(), admin_name="admin")
    d = RegisterUserDialog(app)
    assert d.cb_template.count() == 2
    d.cb_template.setCurrentIndex(1)
    assert d.combos["department"].currentText() == "Бухгалтерия" and d.ou.text() == "OU=Buh,DC=x"
    assert d.template_groups == ["CN=Buh,OU=g"] and "Buh" in d.pv_rows["groups"].text()
    d.cb_template.setCurrentIndex(0)
    assert d.template_groups == [] and d.pv_rows["groups"].text() == "—"
    # чек-лист: кнопка «Создать» активна только когда всё заполнено
    assert not d.btn_create.isEnabled()
    d.surname.setText("Иванов"); d.name.setText("Иван"); d.generate()
    monkeypatch.setattr(config.settings, "use_ssl", True)
    d._refresh()
    assert d.btn_create.isEnabled() and "✅" in d.checks["pwd"].text()
    d.close()


def test_scan_once_non_windows(monkeypatch):
    from adk.workers import PCScannerWorker
    monkeypatch.setattr(os, "name", "posix")
    conn = _NoConn()
    conn.entries = [_Entry("CN=PC-001,OU=c", name="PC-001"), _Entry("CN=SRV-1,OU=c", name="SRV-1")]
    monkeypatch.setattr(config.settings, "host_pattern", r"^PC-\d+$", raising=False)
    monkeypatch.setattr(PCScannerWorker, "workstation_names", staticmethod(lambda entries: ["PC-001"]))
    assert PCScannerWorker.scan_once(lambda: conn) == 1


def test_subprocess_not_called_for_wake(monkeypatch):
    """wake() шлёт UDP-broadcast, а не запускает внешние программы — проверяем, что сокет открылся и закрылся."""
    sent = []

    class Sock:
        def __init__(self, *a, **k):
            pass

        def setsockopt(self, *a):
            pass

        def sendto(self, data, addr):
            sent.append((len(data), addr))

        def close(self):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            pass
    import socket
    monkeypatch.setattr(socket, "socket", Sock)
    assert nettools.wake("00:1A:2B:3C:4D:5E") is True
    assert sent == [(102, ("255.255.255.255", 9)), (102, ("255.255.255.255", 7))]  # порты 9 и 7 — оба принято слушать


# ------------------------------------------------------------------ 8. роли и безопасность
def test_access_two_rights_by_groups(monkeypatch):
    """GT_Admins (в AD только чтение) может всё с ПК, но не трогает объекты AD; IT-Admins — всё."""
    from adk import access
    monkeypatch.setattr(config.settings, "readonly", False)
    monkeypatch.setattr(config.settings, "readonly_group", "ADK-ReadOnly")
    monkeypatch.setattr(config.settings, "admin_groups", ())
    monkeypatch.setattr(config.settings, "pc_admin_groups", ("GT_Admins", "IT-Admins"))
    monkeypatch.setattr(config.settings, "ad_admin_groups", ("IT-Admins",))
    # состоит только в GT_Admins (регистр не важен)
    pc, ad, why = access.evaluate_rights(["Domain Users", "gt_admins"])
    assert (pc, ad) == (True, False) and "IT-Admins" in why
    # полный администратор
    assert access.evaluate_rights(["IT-Admins"]) == (True, True, "")
    # ни в одной группе — ничего не может
    assert access.evaluate_rights(["Domain Users"])[:2] == (False, False)
    # readonly_group отнимает оба права даже у IT-Admins
    assert access.evaluate_rights(["IT-Admins", "ADK-ReadOnly"])[:2] == (False, False)
    # admin_groups (старый ключ) даёт оба права сразу
    monkeypatch.setattr(config.settings, "admin_groups", ("Domain Admins",))
    assert access.evaluate_rights(["Domain Admins"])[:2] == (True, True)
    # задан только ad_admin_groups — право «ПК» есть у всех
    monkeypatch.setattr(config.settings, "admin_groups", ())
    monkeypatch.setattr(config.settings, "pc_admin_groups", ())
    assert access.evaluate_rights(["Domain Users"])[:2] == (True, False)


def test_access_can_respects_action_class():
    from adk import access
    access.set_rights(pc=True, ad=False, reason="GT_Admins")
    try:
        # с ПК — можно
        assert all(access.can(a) for a in ("restart", "shutdown", "wol", "software", "note_add", "plugin_modifying"))
        # объекты AD — нельзя, включая массовые операции и «Группы как у…»
        assert not any(access.can(a) for a in ("create_user", "disable_user", "reset_password", "smartcard",
                                               "group_add", "modify_user", "bulk_disable", "groups_sync"))
        # не изменяющие — всегда
        assert access.can("ping") and access.can("export") and access.can("health")
        assert not access.is_readonly() and access.is_limited()
        assert "AD" in access.badge_text() and "Active Directory" in access.deny_text("reset_password")
    finally:
        access.reset()
    assert access.can("create_user") and access.badge_text() == ""


def test_main_window_hides_only_ad_buttons(qapp, monkeypatch):
    """Инспектор: у «ПК-администратора» остаются перезагрузка и WoL, пропадают «Новый пользователь» и «Группы как у…»."""
    from adk import access
    from adk.main_window import ADApp
    monkeypatch.setattr(ADApp, "start_scan", lambda self: None)
    monkeypatch.setattr(ADApp, "resolve_access", lambda self: None)
    w = ADApp("CORP\\admin", "pwd")
    w.show(); qapp.processEvents()
    access.set_rights(pc=True, ad=False, reason="GT_Admins")
    try:
        w.apply_access(); qapp.processEvents()
        assert w.lbl_readonly.isVisible() and "AD" in w.lbl_readonly.text()
        # инспектор не на экране (стек показывает дашборд), поэтому смотрим isHidden(), а не isVisible()
        assert not w.action_buttons["power"].isHidden() and "wol" not in w.action_buttons   # WoL — пункт меню «Питание ПК»
        assert w.btn_compare.isHidden()
        new_user_btn = next(b for b, a in w._modifying_buttons if a == "create_user")
        assert new_user_btn.isHidden()
    finally:
        access.reset()
        w.apply_access()
    assert not w.btn_compare.isHidden() and not w.lbl_readonly.isVisible()
    w.close()


def test_ldap_filter_escapes_user_input():
    """Ввод «иван*)(cn=*» не должен ломать фильтр: спецсимволы экранируются, а маски добавляет только наш код."""
    from adk.workers import SearchWorker
    flt = SearchWorker(lambda: None, "иван*)(cn=*")._build_filter()
    assert flt and "\\2a\\29\\28cn=\\2a" in flt          # ( ) * превращены в \28 \29 \2a
    assert flt.count("(") == flt.count(")")              # скобки сбалансированы — фильтр остался корректным


def test_audit_calls_never_receive_password_variables():
    """Ни один вызов db.log_action в коде не передаёт переменную с паролем (проверка по исходникам)."""
    import glob
    import re
    bad = []
    for path in glob.glob(os.path.join(os.path.dirname(__file__), "..", "adk", "*.py")):
        src = open(path, encoding="utf-8").read()
        for m in re.finditer(r"log_action\((.*?)\)\n", src, re.S):
            args = re.sub(r"(\"[^\"]*\"|'[^']*'|f\"[^\"]*\")", "", m.group(1))   # выкидываем строковые литералы
            if re.search(r"\b(pwd|password|passwd|new_password)\b", args):
                bad.append((os.path.basename(path), m.group(1)[:80]))
    assert not bad, bad


def _entry(**attrs):
    """Имитация записи ldap3: entry["displayName"].values → ["Иванов Иван"]."""
    class Attr:
        def __init__(self, v):
            self.values = [v] if v is not None else []

    class Entry(dict):
        pass

    return Entry({k: Attr(v) for k, v in attrs.items()})


def _row(**kw):
    base = {k: "—" for k, _l in workers.INVENTORY_COLUMNS}
    base["comp"] = "Не привязан"
    base.update(kw)
    return base


# ------------------------------------------------------------------ 9. опись ПК
def test_build_inventory_rows(monkeypatch):
    """Строка описи собирается из AD + привязки ПК + инвентаря; человек без ПК помечен «Не привязан»."""
    db.save_computer_for_login("ivanov", "WS-101")
    db.db_execute_with_retry(
        "INSERT INTO pc_inventory (computer_name, ip_address, is_online) VALUES ('WS-101', '10.0.2.11', 1)")
    monkeypatch.setattr(netutils, "get_computer_specs_summary",
                        lambda name: "ОС: Windows 11 Pro\nCPU: Intel i5\nОЗУ: 16 ГБ\nДиски: SSD 500 ГБ")
    monkeypatch.setattr(db, "printers_for_computers", lambda comps: {"WS-101": [{"name": "HP LaserJet", "kind": "network", "ip": "10.0.2.50"}]})
    entries = [
        _entry(displayName="Иванов Иван", sAMAccountName="ivanov", title="Админ", department="ИТ"),
        _entry(displayName="Смирнов Алексей", sAMAccountName="smirnov", title="Инженер", department="ИТ"),
        _entry(displayName="", sAMAccountName="svc_backup"),   # без ФИО — служебная учётка, в опись не идёт
    ]
    rows = workers.build_inventory_rows(entries)
    assert [r["login"] for r in rows] == ["ivanov", "smirnov"]
    r = rows[0]
    assert r["comp"] == "WS-101" and r["ip"] == "10.0.2.11" and r["online"] == "да"
    assert r["os"] == "Windows 11 Pro" and r["printers"].startswith("HP LaserJet")
    assert rows[1]["comp"] == "Не привязан" and rows[1]["os"] == "—"
    assert set(r) >= {k for k, _l in workers.INVENTORY_COLUMNS}


def test_build_inventory_rows_cancel():
    entries = [_entry(displayName=f"Пользователь {i}", sAMAccountName=f"u{i}") for i in range(5)]
    assert workers.build_inventory_rows(entries, cancelled=lambda: True) == []


def test_write_inventory_xlsx(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    rows = [_row(fio="Иванов Иван", login="ivanov", comp="WS-101", os="Windows 11"), _row(fio="Смирнов А.", login="smirnov")]
    path = tmp_path / "opis.xlsx"
    workers.write_inventory_xlsx(str(path), "ООО Пример", rows, ("fio", "comp", "os"))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Опись ПК", "Сводка"]
    ws = wb["Опись ПК"]
    assert [c.value for c in ws[4]] == ["Сотрудник", "Имя ПК", "ОС"]     # шапка — 4-я строка, выше титул
    assert ws.cell(5, 1).value == "Иванов Иван" and ws.cell(6, 2).value == "Не привязан"
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref == "A4:C6"
    summary = {ws2.cell(r, 1).value: ws2.cell(r, 2).value for ws2 in [wb["Сводка"]] for r in range(1, 12)}
    assert summary.get("Сотрудников") == 2 and summary.get("С привязанным ПК") == 1


def test_inventory_dialog_preview_and_columns(qapp):
    """Диалог: предпросмотр заполняет таблицу и плитки, галочки колонок меняют таблицу, кнопка Excel активна."""
    from adk.dialogs import InventoryDialog

    fake_app = SimpleNamespace(get_conn=lambda: None)
    d = InventoryDialog(fake_app, None)
    rows = [_row(fio="Иванов Иван", login="ivanov", comp="WS-101", online="да"), _row(fio="Смирнов А.", login="smirnov")]
    d.show_rows(rows)
    assert d.table.rowCount() == 2
    assert d.table.columnCount() == len(workers.INVENTORY_DEFAULT)
    assert d.btn_go.isEnabled()
    d.checks["login"].setChecked(True)
    assert "Логин" in [d.table.horizontalHeaderItem(i).text() for i in range(d.table.columnCount())]
    d._set_all(False)
    assert not d.btn_go.isEnabled() and d.table.columnCount() == 0
    d.close()


def test_compare_pc_prefers_live_software_and_falls_back_to_cache(qapp, temp_db, monkeypatch):
    """Сравнение ПК: список ПО снимается с ПК живьём; если ПК не отвечает — берётся последний сохранённый список,
    и в шапке написано «недоступен, данные от <дата>»."""
    from adk import software
    from adk.fleet import ComparePCDialog
    software.cache_software("PC-B", [{"name": "Office", "version": "1", "publisher": "", "installed": ""}])
    monkeypatch.setattr(software, "get_software", lambda host, timeout=60: (
        {"software": [{"name": "Chrome", "version": "1", "publisher": "", "installed": ""}], "updates": []} if host == "PC-A"
        else {"error": "нет сети"}))
    app = SimpleNamespace(get_conn=lambda: None, admin_name="admin")
    d = ComparePCDialog("PC-A", "PC-B", app)
    _spin(qapp)
    names = {d.t_soft.item(r, 0).text() for r in range(d.t_soft.rowCount())}
    assert names == {"Chrome", "Office"}
    assert "PC-A</b>: сейчас" in d.lbl_soft.text() and "PC-B</b>: недоступен, данные от" in d.lbl_soft.text()
    d.close()
