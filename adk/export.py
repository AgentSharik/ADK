"""Экспорт результатов поиска в CSV/XLSX — одинаковые колонки в обоих форматах."""
from __future__ import annotations

import csv

EXPORT_COLUMNS = (("login", "Логин"), ("full_fio", "ФИО"), ("is_disabled", "Учётка"), ("comp", "Имя ПК"), ("ip", "IP"),
                  ("is_online", "Сеть"), ("phone", "Телефон"), ("ip_phone", "IP-тел"), ("mail", "Почта"), ("office", "Кабинет"),
                  ("address", "Адрес"), ("company", "Организация"), ("dept", "Отдел"), ("title", "Должность"),
                  ("last_logon", "Последний вход"), ("printers", "Принтеры"))


def _cell(row: dict, key: str) -> str:
    v = row.get(key)
    if key == "is_disabled":
        return row.get("account_text") or ("Не активна" if v else "Активна")
    if key == "is_online":
        return "В сети" if v else "Не в сети"
    if key == "printers":
        from .netutils import printer_label
        return "; ".join(printer_label(p) for p in (v or []))
    if key == "full_fio":
        return str(v or row.get("fio") or "")
    return "" if v is None else str(v)


def rows_to_table(rows: list[dict]) -> list[list[str]]:
    return [[label for _, label in EXPORT_COLUMNS]] + [[_cell(r, k) for k, _ in EXPORT_COLUMNS] for r in rows]


def export_rows(rows: list[dict], path: str) -> int:
    """Пишет CSV (utf-8-sig, ';') или XLSX по расширению. Возвращает число строк данных."""
    table = rows_to_table(rows)
    _write_table(table, path)
    return len(table) - 1


def _write_table(table: list[list[str]], path: str, sheet: str = "Результаты") -> None:
    if path.lower().endswith(".xlsx"):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        for r in table:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2563EB")
        for i, _ in enumerate(table[0], 1):
            width = max(len(str(r[i - 1])) for r in table) if table else 10
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)
    else:
        with open(path, "w", encoding="utf-8-sig", newline="") as fh:
            csv.writer(fh, delimiter=";").writerows(table)


INVENTORY_COLUMNS = (("computer_name", "Имя ПК"), ("ip_address", "IP"), ("is_online", "Сеть"), ("current_user", "Пользователь"),
                     ("last_logon", "Последний вход"), ("last_seen_online", "Был в сети"), ("last_checked", "Проверен"))


def export_inventory(path: str) -> int:
    """Весь pc_inventory → xlsx/csv (для CLI и кнопки на дашборде)."""
    from . import db
    rows = db.db_execute_with_retry("SELECT computer_name, ip_address, is_online, current_user, last_logon, last_seen_online, last_checked "
                                    "FROM pc_inventory ORDER BY computer_name", fetch="all") or []
    table = [[label for _, label in INVENTORY_COLUMNS]]
    for r in rows:
        table.append([r[0] or "", r[1] or "", "В сети" if r[2] else "Не в сети", r[3] or "", r[4] or "", r[5] or "", r[6] or ""])
    _write_table(table, path)
    return len(rows)
